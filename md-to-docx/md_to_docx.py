# -*- coding: utf-8 -*-
"""把中文 Markdown 转成打印友好的 DOCX（A4、中文字体、表格、等宽代码块）。

用法: python md_to_docx.py 输入.md [输出.docx]
"""
import os
import re
import sys

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = sys.argv[1] if len(sys.argv) > 1 else "input.md"
DST = sys.argv[2] if len(sys.argv) > 2 else os.path.splitext(SRC)[0] + ".docx"

GRAY = RGBColor(0x59, 0x59, 0x59)
LINK = RGBColor(0x1F, 0x4E, 0x9D)

TOKEN_RE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)\s]+\)|<https?://[^>]+>)")
NUM_RE = re.compile(r"^(\d+)\.\s+(.*)$")
HEAD_RE = re.compile(r"^(#+)\s+(.*)$")

doc = Document()

# --- 页面与样式：A4，页边距适中，正文宋体，标题微软雅黑 ---
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.top_margin = sec.bottom_margin = Cm(1.6)
sec.left_margin = sec.right_margin = Cm(1.8)

normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(10.5)
normal.element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")

for name, size in (("Title", 21), ("Heading 1", 15), ("Heading 2", 12.5)):
    st = doc.styles[name]
    st.font.name = "Microsoft YaHei"
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def run_font(run, size, bold=False, color=None, ascii_f="Calibri", ea="SimSun"):
    run.font.size = Pt(size)
    if bold:
        run.font.bold = True
    if color is not None:
        run.font.color.rgb = color
    run.font.name = ascii_f
    run._element.rPr.rFonts.set(qn("w:eastAsia"), ea)


def add_inline(par, text, size=10.5, color=None):
    for tok in TOKEN_RE.split(text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
            inner = tok[2:-2]
            for sub in re.split(r"(<https?://[^>]+>)", inner):
                if not sub:
                    continue
                if sub.startswith("<http"):
                    r = par.add_run(sub[1:-1])
                    run_font(r, size, True, LINK)
                else:
                    r = par.add_run(sub)
                    run_font(r, size, True, color)
        elif tok.startswith("`") and tok.endswith("`") and len(tok) > 2:
            r = par.add_run(tok[1:-1])
            run_font(r, size, color=color, ascii_f="Consolas", ea="SimHei")
        elif tok.startswith("<http"):
            r = par.add_run(tok[1:-1])
            run_font(r, size, color=LINK)
        elif tok.startswith("[") and "](" in tok:
            m = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", tok)
            if m:
                label, url = m.groups()
                r = par.add_run(label + "（" + url + "）")
                run_font(r, size, color=LINK)
        else:
            r = par.add_run(tok)
            run_font(r, size, color=color)


def shade_par(par, fill):
    ppr = par._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    ppr.append(shd)


def shade_cell(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def hr():
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(6)
    ppr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:color"), "BFBFBF")
    pbdr.append(bottom)
    ppr.append(pbdr)


def build_table(rows):
    def cells(r):
        return [c.strip() for c in r.strip().strip("|").split("|")]

    header = cells(rows[0])
    data = [cells(r) for r in rows[1:] if not re.fullmatch(r"[\s:\-|]+", r)]
    table = doc.add_table(rows=1 + len(data), cols=len(header))
    table.style = "Table Grid"
    tblpr = table._tbl.tblPr
    for el in tblpr.findall(qn("w:tblW")):
        tblpr.remove(el)
    w = OxmlElement("w:tblW")
    w.set(qn("w:w"), "5000")
    w.set(qn("w:type"), "pct")
    tblpr.append(w)
    for j, txt in enumerate(header):
        cell = table.cell(0, j)
        shade_cell(cell, "E7E6E6")
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(1)
        add_inline(p, "**" + txt + "**", size=9)
    for ri, row in enumerate(data, start=1):
        for j in range(len(header)):
            cell = table.cell(ri, j)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            add_inline(p, row[j] if j < len(row) else "", size=9)


with open(SRC, encoding="utf-8") as f:
    text = f.read().replace("⚠️", "▲").replace("⭐", "★")
lines = text.splitlines()

# 页脚居中页码
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fld = OxmlElement("w:fldSimple")
fld.set(qn("w:instr"), " PAGE ")
r = OxmlElement("w:r")
t = OxmlElement("w:t")
t.text = "1"
r.append(t)
fld.append(r)
fp._p.append(fld)

i, n = 0, len(lines)
while i < n:
    s = lines[i].strip()
    if s.startswith("```"):
        i += 1
        while i < n and not lines[i].strip().startswith("```"):
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.space_before = Pt(0)
            pf.space_after = Pt(0)
            pf.line_spacing = 1.0
            shade_par(p, "F2F2F2")
            disp = lines[i].replace(" ", "\u00A0") or "\u00A0"
            r = p.add_run(disp)
            run_font(r, 8.5, ascii_f="Consolas", ea="SimHei")
            i += 1
        i += 1
        continue
    if not s:
        i += 1
        continue
    if s == "---":
        hr()
        i += 1
        continue
    m = HEAD_RE.match(s)
    if m:
        level, htext = len(m.group(1)), m.group(2)
        if level == 1:
            p = doc.add_heading(htext, 0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            p = doc.add_heading(htext, level - 1)
        p.paragraph_format.keep_with_next = True
        i += 1
        continue
    if s.startswith(">"):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Cm(0.5)
        pf.space_after = Pt(3)
        add_inline(p, s.lstrip("> ").strip(), size=9.5, color=GRAY)
        i += 1
        continue
    if s.startswith("|"):
        rows = []
        while i < n and lines[i].strip().startswith("|"):
            rows.append(lines[i])
            i += 1
        build_table(rows)
        continue
    if s.startswith("- "):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Cm(0.75)
        pf.first_line_indent = Cm(-0.4)
        pf.space_after = Pt(2)
        r = p.add_run("•  ")
        run_font(r, 10.5)
        add_inline(p, s[2:])
        i += 1
        continue
    m = NUM_RE.match(s)
    if m:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Cm(0.9)
        pf.first_line_indent = Cm(-0.55)
        pf.space_after = Pt(2)
        r = p.add_run(m.group(1) + ".  ")
        run_font(r, 10.5)
        add_inline(p, m.group(2))
        i += 1
        continue
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    add_inline(p, s)
    i += 1

doc.save(DST)

# --- 校验：重新打开确认可读并输出统计 ---
d2 = Document(DST)
print(f"saved: {DST}  paragraphs={len(d2.paragraphs)} tables={len(d2.tables)}")
