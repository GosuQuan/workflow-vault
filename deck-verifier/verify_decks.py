# -*- coding: utf-8 -*-
"""Verify the generated sample decks: slide counts, pictures, key texts."""
import sys
sys.stdout.reconfigure(encoding="utf-8")
from pptx import Presentation

def check(path, want_slides, want_pics, must_contain):
    prs = Presentation(path)
    n = len(prs.slides.__iter__.__self__._sldIdLst) if False else len(prs.slides._sldIdLst)
    pics = 0
    texts = []
    for slide in prs.slides:
        for sh in slide.shapes:
            if sh.shape_type == 13:  # PICTURE
                pics += 1
            if sh.has_text_frame:
                texts.append(sh.text_frame.text)
            if getattr(sh, "has_table", False):
                for row in sh.table.rows:
                    for cell in row.cells:
                        texts.append(cell.text)
    joined = "\n".join(texts)
    print(f"--- {path}")
    print(f"slides: {len(prs.slides._sldIdLst)} (want {want_slides}); pictures: {pics} (want {want_pics})")
    for t in must_contain:
        print(("OK  " if t in joined else "MISS") + " | " + t)

check("AP微积分BC/课件样稿_Unit2-1_基本求导法则.pptx", 16, 1, [
    "Basic Differentiation Rules", "真题节选 Released Exam Excerpt",
    "y′ = 24x (3x² − 5)³", "Class Practice 课堂练习", "2024 College Board",
    "Exam Structure & Calculator", "3.92", "不允许 No",
])
check("AP物理C/课件样稿_Unit2-2_牛顿运动定律.pptx", 17, 2, [
    "Newton's Laws of Motion", "真题节选 Released Exam Excerpt ①",
    "真题节选 Released Exam Excerpt ②", "a = m₂g/(m₁+m₂)",
    "v(τ) = F₀τ/2m", "tanθ = μ_s",
    "Exam Structure & Calculator", "3.50", "21.7%",
])
